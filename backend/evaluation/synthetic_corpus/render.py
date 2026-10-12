"""Rendering of written documents into files.

A draft is first turned into a list of :class:`Block` objects. The canonical
plain text (used for gold quotes) and every file format are produced from the
same blocks, so quotes found in the canonical text are present in the files.
"""

from __future__ import annotations

import io
import logging
import random
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path
from typing import Literal, Optional

from . import fsutil
from .config import PdfConfig, ScanConfig
from .models import CaseLedger, CostNoteView, DocumentDraft, Identity, PlannedDocument
from .text_utils import format_date_de

logger = logging.getLogger(__name__)

COST_MARKER = "[[KOSTENTABELLE]]"
BlockKind = Literal["header", "title", "heading", "para", "bullet", "table", "closing", "email_header", "spacer"]


@dataclass
class Block:
    """One layout element."""

    kind: BlockKind
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)


@dataclass
class EmailHeaders:
    """Resolved email headers."""

    sender: str
    to: list[str]
    cc: list[str]
    subject: str
    sent_at: datetime


def cost_table_rows(view: CostNoteView) -> list[list[str]]:
    """Rows of the cost table, first row is the header."""
    rows = [["Datum", "Leistung", "TP", "Entlohnung", "ES", "ERV", "Summe"]]
    for line in view.lines:
        rows.append([line.service_date, line.description, line.post, line.base, line.einheitssatz, line.erv, line.net])
    rows.append(["", "Summe Entlohnung (netto)", "", "", "", "", view.fees_net])
    rows.append(["", "20 % USt", "", "", "", "", view.ust])
    if view.court_fees != "EUR 0,00":
        rows.append(["", "Pauschalgebühren", "", "", "", "", view.court_fees])
    if view.barauslagen != "EUR 0,00":
        rows.append(["", "Barauslagen", "", "", "", "", view.barauslagen])
    rows.append(["", "Gesamtbetrag", "", "", "", "", view.total])
    return rows


def _person_line(p: Optional[Identity]) -> str:
    if p is None:
        return ""
    name = p.full_name if p.kind == "natural" else p.name
    return f"{name} <{p.email}>" if p.email else name


def email_headers(draft: DocumentDraft, doc: PlannedDocument, ledger: CaseLedger, rng: random.Random) -> EmailHeaders:
    """Resolve sender/recipients to display strings; fall back to the planned author/recipients."""
    meta = draft.email
    sender_id = meta.sender_id if meta and ledger.identity(meta.sender_id) else doc.author_id
    to_ids = [i for i in (meta.recipient_ids if meta else []) if ledger.identity(i)] or doc.recipient_ids
    cc_ids = [i for i in (meta.cc_ids if meta else []) if ledger.identity(i)]
    d = date.fromisoformat(doc.date_iso)
    sent = datetime(d.year, d.month, d.day, rng.randint(7, 18), rng.randint(0, 59))
    return EmailHeaders(
        sender=_person_line(ledger.identity(sender_id)) or sender_id,
        to=[_person_line(ledger.identity(i)) for i in to_ids if ledger.identity(i)],
        cc=[_person_line(ledger.identity(i)) for i in cc_ids],
        subject=(meta.subject if meta else draft.title),
        sent_at=sent,
    )


def _body_blocks(body: str, cost_note: Optional[CostNoteView]) -> tuple[list[Block], bool]:
    blocks: list[Block] = []
    table_done = False
    for raw in body.splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            continue
        if COST_MARKER in stripped:
            before, _, after = stripped.partition(COST_MARKER)
            if before.strip():
                blocks.append(Block("para", before.strip().replace("**", "")))
            if cost_note is not None and not table_done:
                blocks.append(Block("table", rows=cost_table_rows(cost_note)))
                table_done = True
            if after.strip():
                blocks.append(Block("para", after.strip().replace("**", "")))
            continue
        m = re.match(r"^\s{0,3}#{1,6}\s+(.*)$", line)
        if m:
            blocks.append(Block("heading", m.group(1).replace("**", "").strip()))
            continue
        m = re.match(r"^\s*[-*•]\s+(.*)$", line)
        if m:
            blocks.append(Block("bullet", m.group(1).replace("**", "").strip()))
            continue
        blocks.append(Block("para", stripped.replace("**", "")))
    if cost_note is not None and not table_done:
        blocks.append(Block("table", rows=cost_table_rows(cost_note)))
    return blocks, table_done


def build_blocks(
    draft: DocumentDraft, doc: PlannedDocument, ledger: CaseLedger, rng: random.Random
) -> tuple[list[Block], Optional[EmailHeaders]]:
    """Layout blocks for a draft; emails get resolved headers instead of header lines."""
    blocks: list[Block] = []
    headers: Optional[EmailHeaders] = None
    if doc.is_email:
        headers = email_headers(draft, doc, ledger, rng)
        blocks.append(Block("email_header", f"Von: {headers.sender}"))
        blocks.append(Block("email_header", f"An: {'; '.join(headers.to)}"))
        if headers.cc:
            blocks.append(Block("email_header", f"Cc: {'; '.join(headers.cc)}"))
        blocks.append(Block("email_header", f"Gesendet: {format_date_de(doc.date_iso)} {headers.sent_at:%H:%M}"))
        blocks.append(Block("email_header", f"Betreff: {headers.subject}"))
        blocks.append(Block("spacer"))
    else:
        for line in draft.header_lines:
            if line.strip():
                blocks.append(Block("header", line.strip().replace("**", "")))
        if doc.filename_only_fact_id is None and draft.title.strip():
            blocks.append(Block("spacer"))
            blocks.append(Block("title", draft.title.strip().replace("**", "")))
    body, _ = _body_blocks(draft.body_markdown, doc.cost_note)
    if body and body[0].kind in ("para", "heading") and " ".join(body[0].text.split()) == " ".join(draft.title.split()):
        body = body[1:]  # the writer repeated the title as first body line
    blocks.extend(body)
    if draft.closing_lines:
        blocks.append(Block("spacer"))
        blocks.extend(Block("closing", line.strip()) for line in draft.closing_lines if line.strip())
    return blocks, headers


def canonical_text(blocks: list[Block]) -> str:
    """Plain text of the blocks, one element per line, tables as ``a | b | c``."""
    lines: list[str] = []
    for b in blocks:
        if b.kind == "spacer":
            lines.append("")
        elif b.kind == "table":
            lines.extend(" | ".join(c for c in row) for row in b.rows)
        elif b.kind == "bullet":
            lines.append(f"- {b.text}")
        else:
            lines.append(b.text)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip() + "\n"


# ─── Writers per format ──────────────────────────────────────────────────────


def render_docx(blocks: list[Block], path: Path, *, title: str, author: str, doc_date: str) -> None:
    """Write a DOCX file."""
    from docx import Document
    from docx.shared import Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    props = document.core_properties
    props.title = title
    props.author = author
    props.created = datetime.fromisoformat(doc_date)
    for b in blocks:
        if b.kind == "spacer":
            document.add_paragraph("")
        elif b.kind in ("header", "email_header"):
            run = document.add_paragraph().add_run(b.text)
            run.font.size = Pt(9)
        elif b.kind == "title":
            run = document.add_paragraph().add_run(b.text)
            run.bold = True
            run.font.size = Pt(13)
        elif b.kind == "heading":
            document.add_paragraph().add_run(b.text).bold = True
        elif b.kind == "bullet":
            document.add_paragraph(b.text, style="List Bullet")
        elif b.kind == "table":
            table = document.add_table(rows=len(b.rows), cols=len(b.rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(b.rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
                    if r == 0 or row[1] == "Gesamtbetrag":
                        for p in table.cell(r, c).paragraphs:
                            for run in p.runs:
                                run.bold = True
        else:
            document.add_paragraph(b.text)
    buf = io.BytesIO()
    document.save(buf)
    fsutil.write_bytes(path, buf.getvalue())


def _find_fonts(cfg: PdfConfig) -> tuple[str, str]:
    if cfg.font_regular and cfg.font_bold:
        return cfg.font_regular, cfg.font_bold
    candidates = [
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
    ]
    for regular, bold in candidates:
        if Path(regular).is_file() and Path(bold).is_file():
            return regular, bold
    raise RuntimeError("No Unicode TTF font found; set corpus.pdf.font_regular and corpus.pdf.font_bold")


def pdf_bytes(blocks: list[Block], cfg: PdfConfig, *, title: str, author: str, page_numbers: bool = True) -> bytes:
    """Render blocks to PDF bytes with fpdf2."""
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    regular, bold = _find_fonts(cfg)

    class _PDF(FPDF):
        def footer(self) -> None:  # noqa: D401 - fpdf hook
            if page_numbers:
                self.set_y(-12)
                self.set_font("Body", size=8)
                self.cell(0, 6, f"Seite {self.page_no()} von {{nb}}", align="C")

    pdf = _PDF(format="A4")
    pdf.set_margins(20, 18, 20)
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_font("Body", "", regular)
    pdf.add_font("Body", "B", bold)
    pdf.set_title(title)
    pdf.set_author(author)
    pdf.alias_nb_pages()
    pdf.add_page()
    size = cfg.font_size
    line_h = size * 0.5
    for b in blocks:
        if b.kind == "spacer":
            pdf.ln(line_h)
        elif b.kind in ("header", "email_header"):
            pdf.set_font("Body", size=size - 1.5)
            pdf.multi_cell(0, line_h * 0.9, b.text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        elif b.kind == "title":
            pdf.ln(line_h * 0.5)
            pdf.set_font("Body", "B", size + 2.5)
            pdf.multi_cell(0, line_h * 1.2, b.text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(line_h * 0.5)
        elif b.kind == "heading":
            pdf.ln(line_h * 0.4)
            pdf.set_font("Body", "B", size)
            pdf.multi_cell(0, line_h, b.text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        elif b.kind == "bullet":
            pdf.set_font("Body", size=size)
            pdf.multi_cell(0, line_h, f"\u2022 {b.text}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        elif b.kind == "table":
            pdf.set_font("Body", size=size - 2.5)
            widths = (16, 54, 10, 22, 20, 14, 24)
            with pdf.table(col_widths=widths, text_align=("LEFT", "LEFT", "LEFT", "RIGHT", "RIGHT", "RIGHT", "RIGHT"),
                           line_height=line_h * 0.9, first_row_as_headings=True) as table:
                for row in b.rows:
                    r = table.row()
                    for cell in row:
                        r.cell(cell)
            pdf.ln(line_h * 0.5)
        else:
            pdf.set_font("Body", size=size)
            pdf.multi_cell(0, line_h, b.text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(line_h * 0.25)
    return bytes(pdf.output())


def render_pdf(blocks: list[Block], path: Path, cfg: PdfConfig, *, title: str, author: str) -> int:
    """Write a PDF file and return the page count."""
    data = pdf_bytes(blocks, cfg, title=title, author=author)
    fsutil.write_bytes(path, data)
    import pypdfium2 as pdfium

    return len(pdfium.PdfDocument(data))


def render_txt(blocks: list[Block], path: Path) -> None:
    """Write a UTF-8 text file (internal notes)."""
    fsutil.write_text(path, canonical_text(blocks))


def render_eml(blocks: list[Block], headers: EmailHeaders, path: Path) -> None:
    """Write an RFC 5322 email file; the body is the text after the header blocks."""
    msg = EmailMessage()
    msg["From"] = headers.sender
    msg["To"] = ", ".join(headers.to) or headers.sender
    if headers.cc:
        msg["Cc"] = ", ".join(headers.cc)
    msg["Subject"] = headers.subject
    msg["Date"] = format_datetime(headers.sent_at.astimezone())
    body_blocks = [b for b in blocks if b.kind != "email_header"]
    msg.set_content(canonical_text(body_blocks), charset="utf-8")
    fsutil.write_bytes(path, bytes(msg))


def render_scan(
    blocks: list[Block], stem_path: Path, pdf_cfg: PdfConfig, scan_cfg: ScanConfig, rng: random.Random, *, title: str
) -> list[tuple[Path, str]]:
    """Simulate a scanned paper document as JPEG page images.

    Returns:
        ``(path, page_text)`` per written page; ``page_text`` is the text layer of
        the clean page before degradation, used to map facts to page files.
    """
    import pypdfium2 as pdfium
    from PIL import Image, ImageFilter

    data = pdf_bytes(blocks, pdf_cfg, title=title, author="", page_numbers=False)
    pdf = pdfium.PdfDocument(data)
    n = min(len(pdf), scan_cfg.max_pages)
    out: list[tuple[Path, str]] = []
    for i in range(n):
        page = pdf[i]
        text = page.get_textpage().get_text_bounded()
        img = page.render(scale=scan_cfg.dpi / 72).to_pil().convert("L")
        angle = rng.uniform(-scan_cfg.max_rotation_deg, scan_cfg.max_rotation_deg)
        img = img.rotate(angle, resample=Image.BICUBIC, expand=False, fillcolor=255)
        if scan_cfg.noise > 0:
            noise = Image.effect_noise(img.size, 255 * scan_cfg.noise * 4).convert("L")
            img = Image.blend(img, noise, scan_cfg.noise)
        if scan_cfg.blur_radius > 0:
            img = img.filter(ImageFilter.GaussianBlur(scan_cfg.blur_radius))
        target = stem_path.with_suffix(".jpg") if n == 1 else stem_path.with_name(f"{stem_path.name} - Seite {i + 1}.jpg")
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=scan_cfg.jpeg_quality)
        fsutil.write_bytes(target, buf.getvalue())
        out.append((target, text))
    if len(pdf) > n:
        logger.info("%s: scan truncated to %d of %d pages", stem_path.name, n, len(pdf))
    return out


def soffice_path() -> Optional[str]:
    """Path of LibreOffice's ``soffice`` if installed."""
    found = shutil.which("soffice") or shutil.which("soffice.exe")
    if found:
        return found
    for candidate in ("C:/Program Files/LibreOffice/program/soffice.exe", "/usr/bin/soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(candidate).is_file():
            return candidate
    return None


def convert_docx_to_doc(docx_path: Path, doc_path: Path, timeout_s: float = 120) -> bool:
    """Convert DOCX to legacy DOC with LibreOffice. Returns False if unavailable or failed."""
    exe = soffice_path()
    if not exe:
        return False
    with tempfile.TemporaryDirectory() as tmp:
        try:
            subprocess.run(
                [exe, "--headless", "--convert-to", "doc:MS Word 97", "--outdir", tmp, str(docx_path)],
                check=True, capture_output=True, timeout=timeout_s,
            )
        except (subprocess.SubprocessError, OSError) as exc:
            logger.warning("LibreOffice conversion failed for %s: %s", docx_path.name, exc)
            return False
        produced = Path(tmp) / (docx_path.stem + ".doc")
        if not produced.is_file():
            return False
        fsutil.move(produced, doc_path)
    return True


# ─── Text extraction for validation ──────────────────────────────────────────


def extract_text(path: Path) -> Optional[str]:
    """Extract text from a rendered file the way a simple ingester would (None for images)."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document

        d = Document(io.BytesIO(fsutil.read_bytes(path)))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if suffix == ".pdf":
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(fsutil.read_bytes(path))
        return "\n".join(pdf[i].get_textpage().get_text_bounded() for i in range(len(pdf)))
    if suffix == ".eml":
        from email import policy
        from email.parser import BytesParser

        msg = BytesParser(policy=policy.default).parse(io.BytesIO(fsutil.read_bytes(path)))
        body = msg.get_body(preferencelist=("plain",))
        head = "\n".join(f"{k}: {msg[k]}" for k in ("From", "To", "Cc", "Subject", "Date") if msg[k])
        return head + "\n\n" + (body.get_content() if body else "")
    if suffix == ".txt":
        return fsutil.read_bytes(path).decode("utf-8")
    return None


def random_time_offset(rng: random.Random) -> timedelta:
    """Small helper used for file modification times."""
    return timedelta(hours=rng.randint(8, 18), minutes=rng.randint(0, 59))
